#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Aggregazione indici clinici da piu articoli + motore di valutazione della pericolosita.
Articolo 1: indici cardiologici.  Articolo 2: indici antropometrici."""

import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "Indici_Clinici_Analisi_Rischio.xlsx")
FONT = "Arial"
BLUE, BLACK, GREEN, YEL = "0000FF", "000000", "008000", "FFFF00"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
BAD_FILL = PatternFill("solid", fgColor="F8CBAD")
ERR_FILL = PatternFill("solid", fgColor="FCE4D6")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

A1 = "Art.1 - Indici cardiologici"
A2 = "Art.2 - Indici antropometrici"


def hdr(ws, row, labels, start=1, h=32):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=lab)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = h


def title(ws, text, sub=None):
    c = ws.cell(row=1, column=1, value=text)
    c.font = Font(name=FONT, size=14, bold=True, color="1F3864")
    if sub:
        c2 = ws.cell(row=2, column=1, value=sub)
        c2.font = Font(name=FONT, size=9, italic=True, color="595959")


def setw(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def rowh(nota, width=78, lo=30, hi=80):
    return max(lo, min(hi, 13 * (len(nota) // width + 1)))


wb = Workbook()

# =====================================================================
# 00 LEGENDA
# =====================================================================
ws = wb.active
ws.title = "00 Legenda"
setw(ws, {"A": 32, "B": 100})
title(ws, "Indici clinici - aggregazione multi-articolo, calcolo e valutazione della pericolosita",
      "Fonti attualmente integrate: Art.1 indici cardiologici, Art.2 indici antropometrici. Stesso soggetto dimostrativo.")

LEG = [
    ("", ""),
    ("SCOPO", "Raccogliere in un unico modello formule e indici estratti dagli articoli forniti, ricalcolarli in "
              "modo trasparente dai dati grezzi del soggetto, verificare la riproducibilita dei valori pubblicati "
              "e produrre una valutazione strutturata del livello di pericolosita."),
    ("", ""),
    ("AVVERTENZA", "Strumento didattico e di sintesi documentale. NON e un dispositivo medico e non produce diagnosi. "
                   "'Gennarino Parsifallo' e dichiarato dalla fonte stessa come caso dimostrativo, non un paziente "
                   "reale. Ogni decisione clinica compete al medico."),
    ("", ""),
    ("COME SI USA", "1) Apri '01 Input' e modifica SOLO le celle blu su sfondo giallo.\n"
                    "2) Tutti gli altri fogli si ricalcolano da soli.\n"
                    "3) '02 Indici' contiene i 32 indici delle due fonti, ricalcolati e confrontati con il referto.\n"
                    "4) '06 Rischio' produce il punteggio complessivo di pericolosita.\n"
                    "5) '08 Controlli' elenca cosa e riproducibile e cosa no: leggilo prima di fidarti dei numeri."),
    ("", ""),
    ("CONVENZIONI COLORE", ""),
    ("  Testo blu su giallo", "Dato inserito a mano. E l'unica cosa da modificare."),
    ("  Testo nero", "Risultato calcolato nel foglio stesso."),
    ("  Testo verde", "Valore richiamato da un altro foglio."),
    ("  Riga arancione", "Indice fuori intervallo nella direzione critica."),
    ("  Riga gialla", "Indice dentro l'intervallo ma in zona di guardia."),
    ("", ""),
    ("FOGLI", ""),
    ("  01 Input", "Dati grezzi del soggetto, laboratorio e fattori anamnestici. Unico foglio da compilare."),
    ("  02 Indici", "I 32 indici delle due fonti: formula, valore ricalcolato, scostamento dal referto, intervallo "
                    "di riferimento, stato, peso clinico e punteggio."),
    ("  03 Formule", "Dizionario delle formule con derivazione e verifica numerica."),
    ("  04 Zone FC", "Zone di frequenza cardiaca ricavate dalla frequenza massima."),
    ("  05 Evidenze", "Letteratura citata dalle due fonti con PMID e applicabilita al caso."),
    ("  06 Rischio", "Motore di calcolo: punteggio strumentale per fonte, fattori anamnestici, rischio relativo "
                     "epidemiologico e indice composito di pericolosita."),
    ("  07 Scenari", "Simulazioni what-if sugli interventi discussi dalle fonti."),
    ("  08 Controlli", "Verifica di riproducibilita: quali valori pubblicati si riottengono dai dati grezzi e quali no."),
    ("", ""),
    ("NOTA IMPORTANTE", "Il foglio '08 Controlli' segnala sette scostamenti fra i valori pubblicati e il ricalcolo, "
                        "fra cui un BMI che non corrisponde al peso e all'altezza dichiarati e che si propaga su tre "
                        "indici derivati. Il modello mantiene sia il valore pubblicato sia quello ricalcolato, "
                        "senza sceglierne uno."),
    ("", ""),
    ("AGGIUNGERE UN ARTICOLO", "Aggiungi le righe in fondo alle tabelle di '02 Indici', '03 Formule' e '05 Evidenze' "
                               "compilando la colonna 'Fonte' con l'etichetta del nuovo articolo. Se introduce dati "
                               "grezzi nuovi, inseriscili prima in '01 Input'. I conteggi di '06 Rischio' usano "
                               "COUNTIFS sull'etichetta della fonte, quindi si aggiornano da soli."),
]
r = 4
for a, b in LEG:
    ca, cb = ws.cell(row=r, column=1, value=a), ws.cell(row=r, column=2, value=b)
    ca.font = Font(name=FONT, size=10, bold=bool(a and not a.startswith("  ")),
                   color="C00000" if a in ("AVVERTENZA", "NOTA IMPORTANTE") else "000000")
    cb.font = Font(name=FONT, size=10)
    cb.alignment = Alignment(wrap_text=True, vertical="top")
    ca.alignment = Alignment(vertical="top")
    if b:
        ws.row_dimensions[r].height = rowh(b, 100, 15, 62)
    r += 1

# =====================================================================
# 01 INPUT
# =====================================================================
ws = wb.create_sheet("01 Input")
setw(ws, {"A": 46, "B": 16, "C": 14, "D": 74})
title(ws, "01 - Dati grezzi del soggetto",
      "Modifica solo le celle blu su sfondo giallo. Tutto il resto del workbook si ricalcola.")
hdr(ws, 4, ["Parametro", "Valore", "Unita", "Fonte / Nota"])

INPUTS = [
    ("ANAGRAFICA", None, None, None),
    ("Identificativo soggetto", "Gennarino Parsifallo", "", "Caso dimostrativo dichiarato dalle fonti"),
    ("Eta", 50, "anni", A1 + " / " + A2),
    ("Sesso", "M", "M/F", "Determina il modello di FC massima, la stima del grasso corporeo e la formula del VAI"),
    ("", None, None, None),
    ("MISURE ANTROPOMETRICHE", None, None, None),
    ("Peso", 74.0, "kg", A2),
    ("Altezza", 174, "cm", A2),
    ("Circonferenza vita", 97, "cm", A2 + " - a fine espirazione, a meta fra ultima costa e cresta iliaca. Alimenta nove dei dodici indici antropometrici"),
    ("Circonferenza fianchi", 100, "cm", A2),
    ("Circonferenza polpaccio", 34, "cm", A2 + " - surrogato validato della massa muscolare appendicolare (PMID 30312372)"),
    ("Circonferenza polso", 18, "cm", A2 + " - misura stabile per tutta la vita adulta: e il riferimento fisso del morfotipo"),
    ("Peso all'eta di 30 anni", 68.0, "kg", A2 + " - serve a calcolare l'incremento ponderale ventennale"),
    ("Superficie corporea (BSA)", 1.9, "m2", A1 + " - riportata come 'circa 1,9 m2'; denominatore degli indici cardiologici indicizzati"),
    ("", None, None, None),
    ("MISURE CARDIOVASCOLARI", None, None, None),
    ("Pressione sistolica (SBP)", 128, "mmHg", A1),
    ("Pressione diastolica (DBP)", 84, "mmHg", A1),
    ("Frequenza cardiaca a riposo (FC)", 82, "bpm", A1 + " - unico indice cardiologico fuori range del referto"),
    ("Gittata sistolica (SV)", 66, "mL", A1),
    ("Frequenza cardiaca massima da referto", 175, "bpm", A1 + " - modello logistico sesso-specifico (PMID 12752560 / 11722475); i coefficienti non sono pubblicati, quindi il valore si inserisce a mano"),
    ("Pressione venosa centrale assunta (CVP)", 5, "mmHg", "ASSUNZIONE: non riportata dalle fonti. Serve solo alle resistenze vascolari. Con CVP=5 il calcolo restituisce i 2,7 kU*m2 del referto"),
    ("", None, None, None),
    ("LABORATORIO", None, None, None),
    ("Colesterolo HDL", 41, "mg/dL", A2 + " - preso da solo nessun laboratorio lo segnalerebbe. Alimenta CMI e VAI"),
    ("Trigliceridi", 148, "mg/dL", A2 + " - idem. Alimenta LAP, CMI e VAI"),
    ("", None, None, None),
    ("VALORI NON RICALCOLABILI (inseriti dal referto)", None, None, None),
    ("ABSI z-score da referto", 1.52, "z", A2 + " - il valore grezzo di ABSI e ricalcolabile, ma lo z-score richiede le tabelle NHANES per eta e sesso, che non sono pubblicate. Vedi foglio 08"),
    ("", None, None, None),
    ("FATTORI ANAMNESTICI (0 = assente, 1 = presente)", None, None, None),
    ("Sigarette al giorno", 20, "n/die", A1 + " / " + A2),
    ("Anni di abitudine tabagica", 30, "anni", A1 + " / " + A2),
    ("Miosteatosi documentata", 1, "0/1", A2 + " - infiltrazione adiposa del muscolo; non misurabile con l'antropometria di superficie"),
    ("Sospetta apnea ostruttiva del sonno mai indagata", 1, "0/1", "Russamento e pause respiratorie riferite dalla convivente, nessun accertamento eseguito"),
    ("Sedentarieta / assenza di esercizio strutturato", 1, "0/1", A2 + " - calcetto interrotto a 40 anni e mai ripreso, 30.000 km l'anno in auto"),
    ("", None, None, None),
    ("PARAMETRI DI CONFRONTO E SOGLIE", None, None, None),
    ("FC a riposo di riferimento ottimale", 65, "bpm", "Punto medio della fascia 60-69 bpm, minimo di mortalita (PMID 11337213)"),
    ("Soglia FC a riposo ad alto rischio", 80, "bpm", "Soglia oltre la quale il rischio cambia scala (PMID 26598376, PMID 28067310)"),
    ("Limite superiore FC del referto", 76, "bpm", "Estremo alto dell'intervallo del referto; base del calcolo dei battiti in eccesso"),
    ("Soglia vita per sindrome metabolica (maschio europeo)", 101, "cm", A2 + " - criterio antropometrico; il soggetto ne misura 97 e la casella resta vuota"),
]
r = 5
IR = {}
for name, val, unit, note in INPUTS:
    if name == "":
        r += 1
        continue
    if val is None:
        c = ws.cell(row=r, column=1, value=name)
        c.font = Font(name=FONT, size=10, bold=True, color="1F3864")
        for cc in range(1, 5):
            ws.cell(row=r, column=cc).fill = SUB_FILL
            ws.cell(row=r, column=cc).border = BORDER
        r += 1
        continue
    ws.cell(row=r, column=1, value=name).font = Font(name=FONT, size=10)
    vc = ws.cell(row=r, column=2, value=val)
    vc.font = Font(name=FONT, size=10, bold=True, color=BLUE)
    vc.fill = PatternFill("solid", fgColor=YEL)
    vc.alignment = Alignment(horizontal="center")
    uc = ws.cell(row=r, column=3, value=unit)
    uc.font = Font(name=FONT, size=10)
    uc.alignment = Alignment(horizontal="center")
    nc = ws.cell(row=r, column=4, value=note)
    nc.font = Font(name=FONT, size=9, italic=True, color="595959")
    nc.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = rowh(note, 88, 15, 44)
    for cc in range(1, 5):
        ws.cell(row=r, column=cc).border = BORDER
    IR[name] = r
    r += 1

I = lambda k: f"'01 Input'!$B${IR[k]}"
I_AGE, I_SEX = I("Eta"), I("Sesso")
I_W, I_H = I("Peso"), I("Altezza")
I_WC, I_HIP = I("Circonferenza vita"), I("Circonferenza fianchi")
I_CALF, I_WRIST = I("Circonferenza polpaccio"), I("Circonferenza polso")
I_W30 = I("Peso all'eta di 30 anni")
I_BSA = I("Superficie corporea (BSA)")
I_SBP, I_DBP = I("Pressione sistolica (SBP)"), I("Pressione diastolica (DBP)")
I_HR, I_SV = I("Frequenza cardiaca a riposo (FC)"), I("Gittata sistolica (SV)")
I_HRMAX = I("Frequenza cardiaca massima da referto")
I_CVP = I("Pressione venosa centrale assunta (CVP)")
I_HDL, I_TG = I("Colesterolo HDL"), I("Trigliceridi")
I_ABSIZ = I("ABSI z-score da referto")
I_CIG, I_YRS = I("Sigarette al giorno"), I("Anni di abitudine tabagica")
I_MIO = I("Miosteatosi documentata")
I_OSAS = I("Sospetta apnea ostruttiva del sonno mai indagata")
I_SED = I("Sedentarieta / assenza di esercizio strutturato")
I_HROPT = I("FC a riposo di riferimento ottimale")
I_HRTHR = I("Soglia FC a riposo ad alto rischio")
I_HRREF = I("Limite superiore FC del referto")
I_WCMS = I("Soglia vita per sindrome metabolica (maschio europeo)")

ws.cell(row=IR["Pressione venosa centrale assunta (CVP)"], column=2).comment = Comment(
    "Assunzione dell'analista: le fonti non riportano la CVP.\n"
    "Con CVP = 5 mmHg la formula 80*(MAP-CVP)/CI restituisce 2696 dyn*s*cm-5*m2 = 2,7 kU*m2,\n"
    "cioe il valore stampato sul referto. Con CVP = 0 si otterrebbe 2,8.", "Analisi")
ws.cell(row=IR["Peso"], column=2).comment = Comment(
    "ATTENZIONE: 74 kg su 174 cm danno un BMI di 24,44, non i 24,09 stampati sul referto.\n"
    "Il referto usa internamente 24,09 (lo si deduce dal VAI, che torna solo con quel valore).\n"
    "Vedi foglio '08 Controlli', riga 1.", "Analisi")

# =====================================================================
# 02 INDICI
# =====================================================================
ws = wb.create_sheet("02 Indici")
setw(ws, {"A": 5, "B": 15, "C": 36, "D": 34, "E": 13, "F": 12, "G": 13, "H": 11,
          "I": 10, "J": 10, "K": 12, "L": 11, "M": 15, "N": 7, "O": 8, "P": 9,
          "Q": 26, "R": 66})
title(ws, "02 - Indici estratti dalle fonti, ricalcolati e valutati",
      "Colonna E ricalcolata dai dati grezzi del foglio 01. Colonna G riporta il valore pubblicato: la colonna H rende visibile ogni scostamento.")
hdr(ws, 4, ["#", "Gruppo", "Indice", "Formula", "Valore\ncalcolato", "Unita", "Valore\nreferto",
            "Scost.\n%", "Rif.\nmin", "Rif.\nmax", "Direzione\ncritica", "Posiz.\nrange",
            "Stato", "Peso", "Punt.\n0-3", "Contrib.", "Fonte", "Note e interpretazione"])

R0 = 5
BMI = f"E{R0+16}"          # riga del BMI (art.2)
MAP = f"E{R0+2}"
SEXM = f'IF({I_SEX}="M",1,0)'

IND = [
 # ---------------- ARTICOLO 1 : CARDIOLOGICI ----------------
 (A1, "Pressori", "Pressione sistolica", "misurata", f"={I_SBP}", "mmHg", 128, 100, 129, "Alto", 1, "0",
  "Dentro l'intervallo ma a 1 mmHg dal limite superiore."),
 (A1, "Pressori", "Pressione diastolica", "misurata", f"={I_DBP}", "mmHg", 84, 60, 84, "Alto", 1, "0",
  "Esattamente sul limite superiore. La fonte segnala che la soglia contera quando entrera in uno score."),
 (A1, "Pressori", "Pressione arteriosa media (MAP)", "DBP + 0,4 x PP", f"=E{R0+1}+0.4*E{R0+3}", "mmHg", 101, 70, 105, "Alto", 1, "0",
  "La MAP classica (SBP+2xDBP)/3 darebbe 98,7 e non 101: il valore pubblicato si riproduce solo con DBP + 0,4 x PP. La scelta si propaga sul Modified Shock Index, che ha la MAP al denominatore."),
 (A1, "Pressori", "Pressione di pulsazione (PP)", "SBP - DBP", f"=E{R0}-E{R0+1}", "mmHg", 44, 30, 55, "Alto", 2, "0",
  "Marcatore di elasticita dei grossi vasi e di invecchiamento vascolare (PMID 33793325). Dentro il range: da sorvegliare, non da segnalare."),
 (A1, "Dinamici", "Frequenza cardiaca a riposo", "misurata", f"={I_HR}", "bpm", 82, 72, 76, "Alto", 3, "0",
  "Unico indice cardiologico fuori range, liquidato come 'tachicardia lieve'. E anche quello con il supporto epidemiologico piu solido, ed e il motore aritmetico di meta degli altri indici della sezione."),
 (A1, "Dinamici", "Frequenza cardiaca massima stimata", "modello logistico sesso-specifico", f"={I_HRMAX}", "bpm", 175, 158, 192, "Informativo", 0, "0",
  "Stima anagrafica, non misura: non descrive lo stato del soggetto, quindi peso 0. L'intervallo 158-192 e il 90-110% del predetto. Confronto con 220-eta (170) e Tanaka (173) nel foglio 03."),
 (A1, "Dinamici", "Frequenza cardiaca di riserva", "FC max - FC riposo", f"=E{R0+5}-E{R0+4}", "bpm", 93, 70, 120, "Basso", 1, "0",
  "Formalmente normale, ma alta perche la FC massima e alta per eta, non perche il basale sia buono. Con basale 72 salirebbe a 103 senza alcun miglioramento reale. HR 2,8 per mortalita CV in chi non impegna l'80% della riserva (PMID 17446799)."),
 (A1, "Dinamici", "Modified Shock Index (MSI)", "FC riposo / MAP", f"=E{R0+4}/E{R0+2}", "-", 0.81, 0.70, 0.99, "Alto", 1, "0.00",
  "Va calcolato sulla pressione MEDIA, non sulla sistolica: con la sistolica verrebbe 0,64, errore diffuso nei calcolatori online."),
 (A1, "Dinamici", "Prodotto cardiovascolare (RPP)", "FC riposo x SBP", f"=E{R0+4}*E{R0}", "bpm*mmHg", 10496, 10000, 14999, "Alto", 2, "#,##0",
  "Surrogato non invasivo del consumo miocardico di ossigeno (PMID 38453019). Il soggetto e entrato nella fascia per 496 punti, spinto dalla frequenza. Tre mesi di esercizio supervisionato lo abbassano del 19% (PMID 15793048)."),
 (A1, "Pompa", "Gittata sistolica (SV)", "misurata", f"={I_SV}", "mL", 66, 60, 100, "Basso", 2, "0",
  "Nel 15% piu basso dell'intervallo: e la pompa poco capiente che obbliga il cuore a compensare con la frequenza. E l'indice che spiega perche gli 82 battiti esistono."),
 (A1, "Pompa", "Portata cardiaca (CO)", "SV x FC / 1000", f"=E{R0+9}*E{R0+4}/1000", "L/min", 5.4, 3.5, 6.9, "Basso", 1, "0.0",
  "Il risultato finale e adeguato, ed e per questo che il referto scrive 'normale'. Lo stesso valore si ottiene con SV 80 mL e 68 bpm: identico output, meccanismo diverso."),
 (A1, "Pompa", "Cardiac Index (CI)", "CO / BSA", f"=E{R0+10}/{I_BSA}", "L/min/m2", 2.8, 2.5, 3.9, "Basso", 1, "0.00",
  "Portata indicizzata per superficie corporea. Nel terzo inferiore dell'intervallo."),
 (A1, "Pompa", "Left Cardiac Work Index (LCWI)", "CI x MAP x 2,222", f"=E{R0+11}*E{R0+2}*2.222", "mW/m2", 627, 400, 900, "Bilaterale", 1, "0",
  "La costante 2,222 converte L/min x mmHg in milliwatt (1 L/min = 1,667e-5 m3/s; 1 mmHg = 133,322 Pa). Ricalcolando con il CI non arrotondato si ottiene 639 contro i 627 pubblicati: differenza di solo arrotondamento."),
 (A1, "Pompa", "Indice di rigidita arteriosa", "PP / MAP", f"=E{R0+3}/E{R0+2}", "-", 0.44, 0.30, 0.55, "Alto", 2, "0.00",
  "Rapporto fra componente pulsatile e continua. Ricostruito per riproduzione numerica: PP/MAP da 0,436, mentre PP/SBP darebbe 0,34."),
 (A1, "Pompa", "Resistenze vascolari sistemiche (SVRI)", "80 x (MAP - CVP) / CI / 1000", f"=80*(E{R0+2}-{I_CVP})/E{R0+11}/1000", "kU*m2", 2.7, 1.8, 2.8, "Alto", 2, "0.00",
  "A 0,1 dal limite superiore: le arterie oppongono resistenza vicina al massimo tollerato e il cuore risponde con la frequenza, non aumentando il volume di eiezione. CVP non pubblicata, assunta a 5 mmHg."),
 (A1, "Antropom.", "Superficie corporea (BSA)", "input (formula non pubblicata)", f"={I_BSA}", "m2", 1.9, 1.5, 2.3, "Informativo", 0, "0.00",
  "Inclusa come sedicesima riga: la fonte dichiara 16 indici cardiologici ma ne tabula 15. E il denominatore di CI, LCWI e SVRI. Peso 0."),

 # ---------------- ARTICOLO 2 : ANTROPOMETRICI ----------------
 (A2, "Base", "Indice di massa corporea (BMI)", "peso / altezza^2", f"={I_W}/({I_H}/100)^2", "kg/m2", 24.09, 18.5, 24.9, "Informativo", 0, "0.00",
  "ATTENZIONE: 74 kg su 174 cm danno 24,44, non i 24,09 pubblicati. Il referto usa internamente 24,09 (il VAI torna solo con quel valore). Peso 0 non perche sia irrilevante, ma perche e proprio l'indice che qui assolve un corpo in disordine. Vedi foglio 08."),
 (A2, "Base", "Rapporto altezza / polso", "altezza / circonf. polso", f"={I_H}/{I_WRIST}", "-", 9.7, 9.6, 10.4, "Informativo", 0, "0.00",
  "Morfotipo normolineo. Al polso non c'e ne grasso apprezzabile ne ventre muscolare: e l'unica misura stabile per tutta la vita adulta, e quindi il riferimento fisso contro cui si giudicano gli indici che invece si muovono."),
 (A2, "Base", "Peso ideale secondo Lorenz", "H - 100 - (H - 150)/4", f"={I_H}-100-({I_H}-150)/4", "kg", 68.0, 60, 80, "Informativo", 0, "0.0",
  "Lavora sulla sola statura e ignora eta e corporatura. Estremo basso della forbice del peso desiderabile."),
 (A2, "Base", "Peso ideale secondo Creff", "(H - 100 + eta/10) x 0,9", f"=({I_H}-100+{I_AGE}/10)*0.9", "kg", 71.1, 60, 80, "Informativo", 0, "0.0",
  "Corregge per eta e morfotipo. Con Lorenz definisce la forbice 68-71 kg: il soggetto a 74 kg sfora di meno di tre chili, margine che nessun clinico definirebbe allarmante. Il problema non e quanto pesa ma dove."),
 (A2, "Forma", "Rapporto vita / altezza (WHtR)", "vita / altezza", f"={I_WC}/{I_H}", "-", 0.56, 0.30, 0.52, "Alto", 3, "0.00",
  "La regola aurea vuole la vita sotto la meta dell'altezza. Meta-analisi di 31 studi: discrimina il rischio cardiometabolico meglio del BMI (PMID 22106927). E il primo indice da guardare, per velocita e robustezza."),
 (A2, "Forma", "Rapporto vita / fianchi (WHR)", "vita / fianchi", f"={I_WC}/{I_HIP}", "-", 0.97, 0.70, 0.89, "Alto", 2, "0.00",
  "Forma a mela: distribuzione addominale del grasso, la peggiore dal punto di vista metabolico."),
 (A2, "Forma", "Conicity Index", "(vita/100) / (0,109 x RADQ(peso/altezza_m))", f"=({I_WC}/100)/(0.109*SQRT({I_W}/({I_H}/100)))", "-", 1.36, 1.00, 1.25, "Alto", 2, "0.00",
  "Descrive quanto il tronco si allontana dal cilindro per avvicinarsi al doppio cono. Normalizza la vita per peso e statura, quindi l'addome sporge al netto della corporatura."),
 (A2, "Adiposita", "Percentuale di grasso corporeo", "Deurenberg: 1,2xBMI + 0,23xeta - 10,8xsesso - 5,4", f"=1.2*{BMI}+0.23*{I_AGE}-10.8*{SEXM}-5.4", "%", 23, 8, 25, "Alto", 1, "0.0",
  "Formula non dichiarata dalla fonte: Deurenberg e la piu probabile ma restituisce 24,6 contro i 23 pubblicati. Stima derivata dal BMI, quindi eredita il difetto del BMI di non distinguere i tessuti."),
 (A2, "Adiposita", "Body Adiposity Index (BAI)", "fianchi / altezza_m^1,5 - 18", f"={I_HIP}/(({I_H}/100)^1.5)-18", "%", 25, 8, 23, "Alto", 2, "0.0",
  "Ricava l'adiposita da fianchi e altezza senza mai usare il peso. Quando due indici di adiposita divergono sulla stessa persona vince quello che non usa il peso, perche e il peso a mascherare la ricomposizione corporea."),
 (A2, "Adiposita", "Body Roundness Index (BRI)", "364,2 - 365,5 x RADQ(1 - ((vita/2pi)/(0,5xH))^2)", f"=364.2-365.5*SQRT(1-((({I_WC}/100)/(2*PI()))/(0.5*({I_H}/100)))^2)", "-", 4.50, 1.0, 4.71, "Alto", 1, "0.00",
  "Uno dei due indici assolti del gruppo, e va registrato con onestà. Dentro il riferimento, ma nell'ultimo 10% dell'intervallo."),
 (A2, "Adiposita", "Abdominal Volume Index (AVI)", "(2 x vita^2 + 0,7 x (vita-fianchi)^2) / 1000", f"=(2*{I_WC}^2+0.7*({I_WC}-{I_HIP})^2)/1000", "-", 11.89, 0, 24.5, "Alto", 1, "0.00",
  "NON RIPRODUCIBILE: la formula standard restituisce 18,82 contro gli 11,89 pubblicati. L'esito resta 'dentro il riferimento' con entrambi i valori, quindi la conclusione non cambia, ma il numero non torna. Vedi foglio 08."),
 (A2, "Ibridi", "Lipid Accumulation Product (LAP)", "(vita - 65) x trigliceridi in mmol/L", f"=({I_WC}-65)*({I_TG}/88.57)", "-", 53.5, 0, 26.7, "Alto", 3, "0.0",
  "Il doppio esatto del cut-off. Riconosce il rischio cardiovascolare meglio del BMI (PMID 16150143). Nasce dall'incrocio fra una vita di 97 cm e trigliceridi che presi da soli nessun laboratorio segnalerebbe."),
 (A2, "Ibridi", "Cardiometabolic Index (CMI)", "(vita/altezza) x (TG/HDL) in mmol/L", f"=({I_WC}/{I_H})*(({I_TG}/88.57)/({I_HDL}/38.67))", "-", 0.88, 0, 0.39, "Alto", 2, "0.00",
  "Piu del doppio del cut-off. Buon discriminatore di diabete e disglicemia (PMID 25199852). Va calcolato in mmol/L: usando i mg/dL verrebbe 2,01, un numero che non ha piu alcun rapporto con la soglia."),
 (A2, "Ibridi", "Visceral Adiposity Index (VAI)", "(vita/(39,68+1,88xBMI)) x (TG/1,03) x (1,31/HDL), mmol/L", f"=({I_WC}/(39.68+1.88*{BMI}))*(({I_TG}/88.57)/1.03)*(1.31/({I_HDL}/38.67))", "-", 2.29, 0, 1.92, "Alto", 3, "0.00",
  "Marcatore della FUNZIONE del tessuto adiposo viscerale, non della sua quantita (PMID 20067971). Il referto commenta: disfunzione moderata del tessuto adiposo, resistenza insulinica presente. E l'indice che rivela il BMI usato internamente dal referto."),
 (A2, "Mortalita", "A Body Shape Index (ABSI) - z-score", "z da tabelle NHANES per eta e sesso", f"={I_ABSIZ}", "z", 1.52, -0.272, 0.228, "Alto", 3, "0.00",
  "L'unico indice il cui commento non ammette diplomazia: rischio di mortalita molto alto. Costruito su 14.105 adulti NHANES per isolare la forma corporea depurata da peso e BMI, con aumento quasi esponenziale della mortalita (PMID 22815707). E l'indice che un normopeso non si aspetta mai di trovare alterato. Il valore grezzo di ABSI e ricalcolabile (0,0873), lo z-score no: le tabelle di riferimento non sono pubblicate."),
 (A2, "Muscolo", "Circonferenza polpaccio corretta per BMI", "CC - correzione categoriale EWGSOP2", f"={I_CALF}-IF({BMI}<18.5,-4,IF({BMI}<25,0,IF({BMI}<30,3,7)))", "cm", 33.5, 34.0, 45.0, "Basso", 3, "0.0",
  "Surrogato validato della massa muscolare appendicolare (PMID 30312372). La correzione applicata dal referto non e pubblicata: con quella categoriale EWGSOP2 il BMI del soggetto non comporta alcuna correzione e il valore resterebbe 34,0 cm, cioe ESATTAMENTE sulla soglia invece che sotto. La diagnosi di sarcopenia moderata dipende interamente da quei 5 millimetri non documentati."),
]

r = R0
for i, (fonte, grp, nome, ftxt, fexc, unit, ref, mn, mx, dz, peso, nfmt, nota) in enumerate(IND, 1):
    ws.cell(row=r, column=1, value=i)
    ws.cell(row=r, column=2, value=grp)
    ws.cell(row=r, column=3, value=nome).font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=r, column=4, value=ftxt).font = Font(name=FONT, size=8, italic=True)
    ce = ws.cell(row=r, column=5, value=fexc)
    ce.font = Font(name=FONT, size=10, bold=True, color=BLACK)
    ce.number_format = nfmt
    ws.cell(row=r, column=6, value=unit).font = Font(name=FONT, size=9)
    cg = ws.cell(row=r, column=7, value=ref)
    cg.font = Font(name=FONT, size=10, color=BLUE)
    cg.number_format = nfmt
    ws.cell(row=r, column=8, value=f'=IF(G{r}=0,"",(E{r}-G{r})/ABS(G{r}))').number_format = "0.0%"
    for col, v in ((9, mn), (10, mx)):
        c = ws.cell(row=r, column=col, value=v)
        c.number_format = nfmt
        c.font = Font(name=FONT, size=10, color=BLUE)
    ws.cell(row=r, column=11, value=dz).font = Font(name=FONT, size=9)
    ws.cell(row=r, column=12, value=f'=IF(J{r}=I{r},"",(E{r}-I{r})/(J{r}-I{r}))').number_format = "0%"
    cm = ws.cell(row=r, column=13, value=(
        f'=IF(E{r}>J{r},"Fuori (alto)",IF(E{r}<I{r},"Fuori (basso)",'
        f'IF(AND(K{r}<>"Informativo",OR(AND(OR(K{r}="Alto",K{r}="Bilaterale"),L{r}>0.8),'
        f'AND(OR(K{r}="Basso",K{r}="Bilaterale"),L{r}<0.2))),"Borderline","Normale")))'))
    cm.font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=r, column=14, value=peso).font = Font(name=FONT, size=10, color=BLUE)
    co = ws.cell(row=r, column=15, value=(
        f'=IF(N{r}=0,0,IF(AND(E{r}>J{r},OR(K{r}="Alto",K{r}="Bilaterale")),IF((E{r}-J{r})/(J{r}-I{r})>0.1,3,2),'
        f'IF(AND(E{r}<I{r},OR(K{r}="Basso",K{r}="Bilaterale")),IF((I{r}-E{r})/(J{r}-I{r})>0.1,3,2),'
        f'IF(M{r}="Borderline",1,0))))'))
    co.font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=r, column=16, value=f"=N{r}*O{r}")
    ws.cell(row=r, column=17, value=fonte).font = Font(name=FONT, size=8, color="595959")
    cn = ws.cell(row=r, column=18, value=nota)
    cn.font = Font(name=FONT, size=9)
    cn.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = rowh(nota, 82, 30, 92)
    for cc in range(1, 19):
        cell = ws.cell(row=r, column=cc)
        cell.border = BORDER
        if cell.font.name != FONT:
            cell.font = Font(name=FONT, size=10)
        if cc in (1, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=(cc in (3, 4, 18)))
    r += 1
LAST = r - 1

tr = r + 1
ws.cell(row=tr, column=3, value="TOTALE PUNTEGGIO STRUMENTALE (entrambe le fonti)").font = Font(name=FONT, size=10, bold=True)
ws.cell(row=tr, column=14, value=f"=SUM(N{R0}:N{LAST})").font = Font(name=FONT, size=10, bold=True)
ws.cell(row=tr, column=16, value=f"=SUM(P{R0}:P{LAST})").font = Font(name=FONT, size=10, bold=True)
for cc in range(3, 18):
    ws.cell(row=tr, column=cc).fill = SUB_FILL
    ws.cell(row=tr, column=cc).border = BORDER

ws.freeze_panes = "E5"
ws.auto_filter.ref = f"A4:R{LAST}"
rng = f"A{R0}:R{LAST}"
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($M{R0},5)="Fuori"'], fill=BAD_FILL))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$M{R0}="Borderline"'], fill=WARN_FILL))
ws.conditional_formatting.add(f"H{R0}:H{LAST}",
                              FormulaRule(formula=[f'ABS($H{R0})>0.02'], fill=ERR_FILL,
                                          font=Font(name=FONT, bold=True, color="C00000")))

C = lambda off: f"'02 Indici'!$E${R0+off}"
E_WHtR, E_LAP, E_CMI, E_VAI, E_CC = C(20), C(27), C(28), C(29), C(31)
E_BMI = f"'02 Indici'!${BMI}"

# =====================================================================
# 03 FORMULE
# =====================================================================
ws = wb.create_sheet("03 Formule")
setw(ws, {"A": 5, "B": 13, "C": 32, "D": 44, "E": 26, "F": 18, "G": 66})
title(ws, "03 - Dizionario delle formule",
      "Le formule marcate 'Ricostruita' non sono scritte per esteso nelle fonti: sono state ricavate cercando l'espressione che riproduce il valore pubblicato dai dati grezzi noti.")
hdr(ws, 4, ["#", "Fonte", "Grandezza", "Formula", "Verifica sul caso", "Origine", "Nota metodologica"])

FORM = [
 (A1, "Pressione arteriosa media", "MAP = DBP + 0,4 x (SBP - DBP)", "84 + 0,4x44 = 101,6 -> 101", "Ricostruita",
  "La formula classica (SBP + 2xDBP)/3 restituisce 98,7 e NON riproduce il valore pubblicato. Il coefficiente 0,4 al posto di 1/3 tiene conto dell'accorciamento della diastole. Chi ricalcola con la classica ottiene un Modified Shock Index di 0,83 invece di 0,81."),
 (A1, "Pressione di pulsazione", "PP = SBP - DBP", "128 - 84 = 44", "Esplicita",
  "Cresce quando l'aorta perde elasticita e smette di ammortizzare l'onda sistolica."),
 (A1, "FC massima - formula popolare", "FCmax = 220 - eta", "220 - 50 = 170", "Citata come errata",
  "Sottostima la FC massima negli adulti maturi. E la formula che fa sospettare un errore a chi verifica il referto."),
 (A1, "FC massima - Tanaka", "FCmax = 208 - 0,7 x eta", "208 - 35 = 173", "PMID 11153730",
  "Meta-analisi su 351 studi e 18.712 soggetti. Migliora la stima ma resta lineare, quindi applica lo stesso decadimento a ogni eta."),
 (A1, "FC massima - modello del referto", "modello logistico sesso-specifico", "175 (da referto)", "PMID 12752560 / 11722475",
  "Tarato su cicloergometro fino a esaurimento. Decadimento non costante: 0,43 bpm/anno a 25 anni, 0,83 a 50, 1,35 a 75. Coefficienti non pubblicati."),
 (A1, "FC di riserva", "FCris = FCmax - FCriposo", "175 - 82 = 93", "Esplicita",
  "Migliora abbassando il basale, cioe senza che il cuore cambi davvero: e un artefatto da tenere presente."),
 (A1, "Modified Shock Index", "MSI = FCriposo / MAP", "82 / 101 = 0,81", "Esplicita",
  "La fonte segnala l'errore diffuso di usare la sistolica al denominatore: 82/128 darebbe 0,64."),
 (A1, "Prodotto cardiovascolare", "RPP = FCriposo x SBP", "82 x 128 = 10.496", "Esplicita / PMID 38453019",
  "Rate pressure product, surrogato del consumo miocardico di ossigeno."),
 (A1, "Portata cardiaca", "CO = SV x FC", "66 x 82 = 5.412 mL/min", "Esplicita",
  "Lo stesso output si ottiene con SV 80 mL e FC 68: la portata normale non distingue una pompa capiente lenta da una piccola veloce."),
 (A1, "Cardiac Index", "CI = CO / BSA", "5,412 / 1,9 = 2,85 -> 2,8", "Esplicita", "Portata indicizzata per superficie corporea."),
 (A1, "Left Cardiac Work Index", "LCWI = CI x MAP x 2,222", "2,8 x 101 x 2,222 = 628 (referto 627)", "Ricostruita",
  "La costante converte L/min x mmHg in milliwatt: 1 L/min = 1,6667e-5 m3/s e 1 mmHg = 133,322 Pa, prodotto 2,222 mW. Con il CI non arrotondato si ottiene 639, circa il 2% sopra il pubblicato."),
 (A1, "Indice di rigidita arteriosa", "ASI = PP / MAP", "44 / 101 = 0,436 -> 0,44", "Ricostruita",
  "Il rapporto alternativo PP/SBP darebbe 0,34 e non riproduce il referto."),
 (A1, "Resistenze vascolari sistemiche", "SVRI = 80 x (MAP - CVP) / CI", "80 x (101-5) / 2,85 = 2.696 -> 2,7 kU*m2", "Ricostruita",
  "Formula standard in dyn*s*cm-5*m2, espressa dal referto in kilo-unita. CVP non pubblicata, assunta a 5 mmHg: con CVP 0 si otterrebbe 2,8."),
 (A1, "Zone di allenamento", "inizio zona = INT(% x FCmax)", "INT(0,5x175)=87; INT(0,6x175)=105", "Ricostruita",
  "Le soglie del referto si riproducono con il troncamento all'intero, non con l'arrotondamento: 0,5x175 = 87,5 diventa 87 e non 88."),
 (A1, "Battiti in eccesso all'anno", "(FC - limite sup.) x 60 x 24 x 365", "6 x 525.600 = 3.153.600", "Esplicita",
  "Traduce lo scostamento istantaneo in carico cumulativo annuo."),
 (A1, "Rischio relativo da FC a riposo", "RR = 1,09^((FC - FCrif)/10)", "1,09^((82-65)/10) = 1,16", "PMID 26598376",
  "Estrapolazione log-lineare dell'incremento per ogni 10 bpm. Stima di gradiente, non rischio assoluto individuale."),
 (A2, "Indice di massa corporea", "BMI = peso / altezza_m^2", "74 / 1,74^2 = 24,44 (referto 24,09)", "Esplicita",
  "SCOSTAMENTO: il valore pubblicato non e ottenibile dal peso e dall'altezza dichiarati. Servirebbero 72,9 kg a 174 cm, oppure 175,3 cm a 74 kg. Il BMI entra in VAI, ABSI e nella correzione del polpaccio."),
 (A2, "Morfotipo", "altezza / circonferenza polso", "174 / 18 = 9,67 -> 9,7", "Esplicita",
  "Normolineo se compreso fra 9,6 e 10,4. Al polso non c'e ne grasso apprezzabile ne ventre muscolare: e l'unica misura che non cambia con dieta o allenamento."),
 (A2, "Peso ideale - Lorenz", "P = H - 100 - (H - 150)/4", "174 - 100 - 6 = 68,0", "Ricostruita",
  "Riproduce esattamente il valore pubblicato. Usa la sola statura e ignora eta e corporatura."),
 (A2, "Peso ideale - Creff", "P = (H - 100 + eta/10) x 0,9", "(174-100+5) x 0,9 = 71,1", "Ricostruita",
  "Riproduce esattamente il valore pubblicato. Corregge per eta e morfotipo: il peso ottimale di un cinquantenne normolineo non e quello di un venticinquenne longilineo."),
 (A2, "Rapporto vita / altezza", "WHtR = vita / altezza", "97 / 174 = 0,557 -> 0,56", "Esplicita / PMID 22106927",
  "Meta-analisi di 31 studi: discrimina il rischio cardiometabolico meglio del BMI. Se supera 0,50 vale la pena calcolare tutta la batteria."),
 (A2, "Rapporto vita / fianchi", "WHR = vita / fianchi", "97 / 100 = 0,97", "Esplicita", "Cut-off 0,89 per il maschio. Forma a mela."),
 (A2, "Conicity Index", "CI = (vita_m) / (0,109 x RADQ(peso/altezza_m))", "0,97 / (0,109 x 6,522) = 1,3646 -> 1,36", "Ricostruita",
  "Riproduce esattamente il valore pubblicato. Misura quanto il tronco si allontana dal cilindro per avvicinarsi al doppio cono, al netto di peso e statura."),
 (A2, "Percentuale di grasso corporeo", "Deurenberg: 1,20xBMI + 0,23xeta - 10,8xsesso - 5,4", "24,6 con BMI ricalcolato; 24,2 con BMI 24,09 (referto 23)", "Non riproducibile",
  "La fonte non dichiara la formula. Deurenberg e la candidata piu probabile ma non restituisce il valore pubblicato con nessuno dei due BMI. Scostamento di circa 1,2 punti percentuali."),
 (A2, "Body Adiposity Index", "BAI = fianchi / altezza_m^1,5 - 18", "100 / 2,2952 - 18 = 25,57 -> 25", "Ricostruita",
  "Riproduce il valore pubblicato. Non usa mai il peso: e per questo che quando diverge dalla stima basata sul BMI, la fonte da ragione al BAI."),
 (A2, "Body Roundness Index", "BRI = 364,2 - 365,5 x RADQ(1 - ((vita/2pi)/(0,5xH))^2)", "364,2 - 359,70 = 4,5005 -> 4,50", "Ricostruita",
  "Riproduce esattamente il valore pubblicato, alla seconda cifra decimale."),
 (A2, "Abdominal Volume Index", "AVI = (2 x vita^2 + 0,7 x (vita-fianchi)^2) / 1000", "18,824 (referto 11,89)", "Non riproducibile",
  "SCOSTAMENTO del 58%. La formula standard di Guerrero-Romero e Rodriguez-Moran non restituisce il valore pubblicato. L'esito qualitativo non cambia (dentro il riferimento con entrambi i valori), ma il numero non torna."),
 (A2, "Lipid Accumulation Product", "LAP = (vita - 65) x TG in mmol/L", "32 x 1,671 = 53,47 -> 53,5", "Ricostruita / PMID 16150143",
  "Riproduce il valore pubblicato. I trigliceridi vanno convertiti in mmol/L dividendo i mg/dL per 88,57."),
 (A2, "Cardiometabolic Index", "CMI = (vita/altezza) x (TG/HDL) in mmol/L", "0,5575 x 1,576 = 0,879 -> 0,88", "Ricostruita / PMID 25199852",
  "Riproduce il valore pubblicato SOLO in mmol/L. Con i mg/dL verrebbe 2,01, che non ha piu alcun rapporto con il cut-off di 0,39."),
 (A2, "Visceral Adiposity Index", "VAI(M) = (vita/(39,68+1,88xBMI)) x (TG/1,03) x (1,31/HDL)", "2,2883 con BMI 24,09 -> 2,29; 2,2706 con BMI 24,44", "Ricostruita / PMID 20067971",
  "Riproduce il valore pubblicato solo usando il BMI di 24,09: e la prova che il referto usa internamente quel BMI e non quello ricalcolabile da peso e altezza. TG e HDL in mmol/L."),
 (A2, "A Body Shape Index", "ABSI = vita_m / (BMI^(2/3) x RADQ(altezza_m))", "0,97 / 11,11 = 0,0873", "Parzialmente ricostruita / PMID 22815707",
  "Il valore grezzo si ricalcola. Lo z-score pubblicato (1,52) richiede le tabelle NHANES per eta e sesso, non riportate dalla fonte, quindi nel foglio 02 viene inserito a mano."),
 (A2, "Circonferenza polpaccio corretta", "CC - correzione categoriale per BMI (EWGSOP2)", "34,0 - 0 = 34,0 (referto 33,5)", "Non riproducibile",
  "La correzione categoriale non prevede alcun aggiustamento per BMI fra 18,5 e 24,9, quindi il valore resterebbe 34,0 cm, esattamente sulla soglia. Il referto applica -0,5 cm con una formula non pubblicata, e da quei 5 millimetri dipende la diagnosi di sarcopenia moderata."),
 (A2, "Volume del tessuto adiposo", "a parita di massa il grasso occupa circa il 18% di spazio in piu del muscolo", "-", "Esplicita",
  "Spiega perche la vita cresce mentre l'ago della bilancia sta fermo, e perche il paziente giura in buona fede di non essere ingrassato."),
]
r = 5
for i, (fonte, nome, f, ver, orig, nota) in enumerate(FORM, 1):
    ws.cell(row=r, column=1, value=i).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(row=r, column=2, value=fonte.split(" - ")[0]).font = Font(name=FONT, size=9, color="595959")
    ws.cell(row=r, column=3, value=nome).font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=r, column=4, value=f).font = Font(name=FONT, size=9, color="C00000")
    ws.cell(row=r, column=5, value=ver).font = Font(name=FONT, size=9)
    bad = orig.startswith("Non riproducibile")
    ws.cell(row=r, column=6, value=orig).font = Font(
        name=FONT, size=9, bold=(orig == "Ricostruita" or bad),
        color="C00000" if bad else ("7030A0" if orig == "Ricostruita" else "595959"))
    cn = ws.cell(row=r, column=7, value=nota)
    cn.font = Font(name=FONT, size=9)
    ws.row_dimensions[r].height = rowh(nota, 82, 28, 78)
    for cc in range(1, 8):
        cell = ws.cell(row=r, column=cc)
        cell.border = BORDER
        if cell.font.name != FONT:
            cell.font = Font(name=FONT, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top",
                                   horizontal="center" if cc == 1 else "left")
    if bad:
        for cc in range(1, 8):
            ws.cell(row=r, column=cc).fill = ERR_FILL
    r += 1
ws.freeze_panes = "C5"
ws.auto_filter.ref = f"A4:G{r-1}"

# =====================================================================
# 04 ZONE FC
# =====================================================================
ws = wb.create_sheet("04 Zone FC")
setw(ws, {"A": 8, "B": 26, "C": 12, "D": 12, "E": 12, "F": 14, "G": 62})
title(ws, "04 - Zone di frequenza cardiaca",
      "Ricavate dalla FC massima del foglio 01. Le soglie si ottengono troncando all'intero la percentuale della FC massima.")
hdr(ws, 4, ["Zona", "Obiettivo", "% min", "% max", "Da (bpm)", "A (bpm)", "Nota"])
ZONES = [(0, "Riposo", 0.0, 0.5, "Zona in cui si trova il soggetto a riposo."),
         (1, "Attivita moderata", 0.5, 0.6, "Il limite inferiore e il confine che il soggetto sfiora da fermo."),
         (2, "Controllo del peso", 0.6, 0.7, ""),
         (3, "Aerobiosi", 0.7, 0.8, ""),
         (4, "Anaerobiosi", 0.8, 0.9, ""),
         (5, "Massimale", 0.9, 1.0, "Da leggere con la tolleranza del 90-110% sulla FC massima predetta.")]
r = 5
for z, obj, pmin, pmax, nota in ZONES:
    ws.cell(row=r, column=1, value=z).font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=r, column=2, value=obj)
    ws.cell(row=r, column=3, value=pmin).number_format = "0%"
    ws.cell(row=r, column=4, value=pmax).number_format = "0%"
    ws.cell(row=r, column=5, value=0 if z == 0 else f"=INT(C{r}*{I_HRMAX})")
    ws.cell(row=r, column=6, value=f"={I_HRMAX}" if z == 5 else f"=INT(C{r+1}*{I_HRMAX})-1")
    ws.cell(row=r, column=7, value=nota).font = Font(name=FONT, size=9)
    for cc in range(1, 8):
        cell = ws.cell(row=r, column=cc)
        cell.border = BORDER
        if cell.font.name != FONT:
            cell.font = Font(name=FONT, size=10)
        cell.alignment = Alignment(horizontal="center" if cc in (1, 3, 4, 5, 6) else "left",
                                   vertical="center", wrap_text=(cc == 7))
    ws.row_dimensions[r].height = 26
    r += 1
r += 1
for k, (lab, val, col) in enumerate([
        ("FC a riposo del soggetto", f"={I_HR}", GREEN),
        ("Soglia di uscita dalla zona 0", "=E6", BLACK),
        ("Margine residuo prima di uscire dal riposo", f"=E{r+1}-E{r}", "C00000")]):
    ws.cell(row=r + k, column=2, value=lab).font = Font(name=FONT, size=10, bold=True)
    c = ws.cell(row=r + k, column=5, value=val)
    c.font = Font(name=FONT, size=11 if k == 2 else 10, bold=True, color=col)
    c.alignment = Alignment(horizontal="center")
    for cc in range(2, 8):
        ws.cell(row=r + k, column=cc).border = BORDER
ws.cell(row=r + 2, column=7, value="Battiti che separano il soggetto, da fermo, dalla zona dell'attivita moderata.").font = Font(name=FONT, size=9, italic=True)

# =====================================================================
# 05 EVIDENZE
# =====================================================================
ws = wb.create_sheet("05 Evidenze")
setw(ws, {"A": 5, "B": 11, "C": 12, "D": 34, "E": 24, "F": 26, "G": 16, "H": 22, "I": 56})
title(ws, "05 - Evidenze citate dalle fonti",
      "Ogni riga riporta la misura di rischio come pubblicata e la sua applicabilita al caso in esame.")
hdr(ws, 4, ["#", "Fonte", "PMID", "Studio / disegno", "Numerosita", "Esposizione", "Misura", "Esito", "Applicabilita al caso"])
EV = [
 (A1, "26598376", "Meta-analisi di 46 studi prospettici", "1.246.203 soggetti / 78.349 decessi", "FC a riposo, per ogni +10 bpm", "RR 1,09", "Mortalita totale",
  "Il soggetto e a 82 bpm: 17 bpm oltre la fascia di minimo, circa +16% di rischio relativo per estrapolazione."),
 (A1, "26598376", "Meta-analisi - analisi per categorie", "1.246.203 soggetti", "FC a riposo > 80 bpm", "RR 1,45", "Mortalita totale",
  "SOGLIA SUPERATA (82 bpm). E il salto di scala descritto dalla fonte."),
 (A1, "26598376", "Meta-analisi - analisi per categorie", "1.246.203 soggetti", "FC a riposo > 80 bpm", "RR 1,33", "Mortalita cardiovascolare", "Soglia superata."),
 (A1, "11337213", "Coorte giapponese, 18 anni", "573 uomini, 40-64 anni", "FC > 90 vs 60-69 bpm", "RR 2,68", "Mortalita",
  "Il minimo di mortalita cade fra 60 e 69 bpm e l'aumento comincia ben prima dei 90: il soggetto e nel tratto ascendente."),
 (A1, "28067310", "Coorte di Kailuan, 3 misurazioni in 4 anni", "47.311 adulti", "FC >= 80 bpm in tutte e tre le misurazioni", "HR 1,86", "Mortalita per ogni causa",
  "Lo studio piu pertinente perche misura la persistenza. Applicabile solo se il valore viene confermato nel tempo: servono misurazioni ripetute."),
 (A1, "17446799", "Coorte prospettica, 5 anni", "1.910 uomini", "Riserva impegnata sotto l'80%", "HR 2,8", "Mortalita cardiovascolare",
  "NON VALUTABILE: richiede un test da sforzo che il referto non contiene. E il principale accertamento mancante."),
 (A1, "10353296", "Revisione (Palatini) sull'ipertono simpatico", "-", "Tachicardia cronica", "-", "Meccanismo",
  "Doppio meccanismo: marcatore di iperattivita adrenergica e stimolo meccanico diretto, con stress di parete che favorisce la placca."),
 (A1, "33793325", "Revisione su Circulation Research", "-", "Pressione di pulsazione elevata", "-", "Invecchiamento vascolare",
  "La pulsatilita si scarica su cervello e rene e riduce la diastolica utile alla perfusione coronarica. Il soggetto e dentro il range."),
 (A1, "38453019", "Letteratura sul rate pressure product", "-", "Prodotto cardiovascolare", "-", "Consumo miocardico di O2",
  "Valida il prodotto cardiovascolare come surrogato non invasivo del consumo miocardico di ossigeno."),
 (A1, "35328926", "Effetti acuti del fumo", "giovani normotesi", "Singola sigaretta", "-", "Aumento acuto di MAP e FC",
  "Effetto acuto: su un referto a riposo conta poco. Rileva cio che resta fra una sigaretta e l'altra."),
 (A1, "8689656", "Caso-controllo su variabilita RR", "20 fumatori > 20 sig/die vs controlli", "Fumo abituale", "LF 70,6 vs 46,0; HF 22,1 vs 42,0", "Assetto autonomico spostato",
  "APPLICABILE: il soggetto fuma 20 sigarette al giorno. Piu acceleratore simpatico, meno freno vagale. Parte dei battiti in eccesso viene da qui."),
 (A1, "16029383", "Studio su forti fumatori", "-", "Durata dell'abitudine", "-", "Calo della modulazione vagale",
  "APPLICABILE E PIU PERTINENTE: misura gli anni, non solo le sigarette. Il soggetto ne ha 30."),
 (A1, "15301332", "Letteratura su apnee ostruttive del sonno", "-", "OSAS", "-", "FC a riposo alta, R-R ridotta, variabilita pressoria",
  "SOSPETTO NON INDAGATO: russamento e pause riferite, nessun accertamento. Spiegherebbe parte della frequenza elevata."),
 (A1, "12609010", "Traffico nervoso simpatico", "-", "OSAS", "-", "Simpatico elevato anche in veglia",
  "L'effetto non si esaurisce con il risveglio: persiste con respiro normale e senza ipossia."),
 (A1, "15793048", "Randomizzato controllato", "136 uomini con coronaropatia", "3 mesi di esercizio supervisionato", "-19% RPP", "Prodotto cardiovascolare a riposo",
  "INTERVENTO EFFICACE. Nel gruppo non supervisionato il valore e rimasto invariato: la supervisione e parte dell'effetto. Simulato nel foglio 07."),
 (A1, "8651120", "Coorte di cessazione tabagica", "54 fumatori da >= 1 pacchetto/die", "Cessazione", "-", "Calo FC, miglioramento della variabilita",
  "EFFICACE ma non risolutivo nel breve: a 4 settimane la FC restava sopra e la variabilita sotto i valori attesi per l'eta."),
 (A1, "11153730", "Meta-analisi su FC massima", "351 studi / 18.712 soggetti", "Eta", "FCmax = 208 - 0,7 x eta", "Stima della FC massima",
  "Termine di confronto con il modello logistico del referto."),
 (A2, "22815707", "Coorte NHANES (Krakauer)", "14.105 adulti", "Forma corporea depurata da peso e BMI", "aumento quasi esponenziale", "Mortalita",
  "APPLICABILE E IL PIU GRAVE: z-score 1,52 contro un limite di 0,228. E l'indice che un normopeso non si aspetta mai di trovare alterato."),
 (A2, "22106927", "Meta-analisi (Ashwell)", "31 studi", "Rapporto vita/altezza", "-", "Rischio cardiometabolico",
  "APPLICABILE: WHtR 0,56 contro 0,52. Il rapporto vita/altezza discrimina meglio del BMI. Un metro da sarta batte la bilancia."),
 (A2, "16150143", "Studio di validazione (Kahn)", "-", "Lipid Accumulation Product", "-", "Rischio cardiovascolare",
  "APPLICABILE: LAP 53,5 contro 26,7, il doppio esatto del cut-off. Riconosce il rischio meglio del BMI."),
 (A2, "20067971", "Validazione (Amato)", "-", "Visceral Adiposity Index", "-", "Funzione del tessuto adiposo",
  "APPLICABILE: VAI 2,29 contro 1,92. Marcatore della funzione del grasso viscerale, non della sua quantita."),
 (A2, "25199852", "Costruzione dell'indice (Wakabayashi e Daimon)", "-", "Cardiometabolic Index", "-", "Diabete e disglicemia",
  "APPLICABILE: CMI 0,88 contro 0,39, piu del doppio. Buon discriminatore di disglicemia."),
 (A2, "26551006", "Coorte (Sahakyan)", "-", "Obesita centrale in soggetti normopeso", "-", "Mortalita cardiovascolare",
  "DESCRIVE ESATTAMENTE IL CASO: nel normopeso l'obesita centrale si accompagna a mortalita superiore a quella dell'obeso con distribuzione periferica. Il magro con il girovita largo e il caso peggiore, non il migliore."),
 (A2, "30312372", "Consenso europeo EWGSOP2", "-", "Circonferenza del polpaccio", "-", "Massa muscolare appendicolare",
  "APPLICABILE: valida il polpaccio come surrogato della massa muscolare. La correzione per BMI usata dal referto non e pero quella del consenso: vedi foglio 08."),
 (A2, "35227529", "Consenso ESPEN-EASO (Donini)", "-", "Obesita sarcopenica", "-", "Insulino-resistenza, diabete 2, ridotta mobilita",
  "APPLICABILE: il soggetto non e obeso secondo il BMI ma ha gia il fenotipo. Formalizza l'associazione fra infiltrazione adiposa del muscolo e insulino-resistenza."),
]
r = 5
for i, (fonte, pmid, studio, n, esp, mis, esito, appl) in enumerate(EV, 1):
    ws.cell(row=r, column=1, value=i)
    ws.cell(row=r, column=2, value=fonte.split(" - ")[0]).font = Font(name=FONT, size=9, color="595959")
    cp = ws.cell(row=r, column=3, value=pmid)
    cp.font = Font(name=FONT, size=10, color="0563C1", underline="single")
    cp.hyperlink = f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/"
    ws.cell(row=r, column=4, value=studio)
    ws.cell(row=r, column=5, value=n)
    ws.cell(row=r, column=6, value=esp)
    ws.cell(row=r, column=7, value=mis).font = Font(name=FONT, size=9, bold=True, color="C00000")
    ws.cell(row=r, column=8, value=esito)
    ws.cell(row=r, column=9, value=appl)
    for cc in range(1, 10):
        cell = ws.cell(row=r, column=cc)
        cell.border = BORDER
        if cell.font.name != FONT:
            cell.font = Font(name=FONT, size=9)
        cell.alignment = Alignment(wrap_text=True, vertical="top",
                                   horizontal="center" if cc in (1, 2, 3, 7) else "left")
    ws.row_dimensions[r].height = rowh(appl, 58, 28, 70)
    r += 1
ws.freeze_panes = "D5"
ws.auto_filter.ref = f"A4:I{r-1}"

# =====================================================================
# 06 RISCHIO
# =====================================================================
ws = wb.create_sheet("06 Rischio")
setw(ws, {"A": 4, "B": 54, "C": 16, "D": 12, "E": 12, "F": 10, "G": 66})
title(ws, "06 - Valutazione della pericolosita",
      "Tre letture indipendenti: il punteggio strumentale dagli indici delle due fonti, i fattori anamnestici che i referti non contengono e il rischio relativo epidemiologico.")

def section(row, text):
    ws.cell(row=row, column=2, value=text).font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
    for cc in range(2, 8):
        ws.cell(row=row, column=cc).fill = HDR_FILL
        ws.cell(row=row, column=cc).border = BORDER
    ws.row_dimensions[row].height = 20

def line(row, label, formula, fmt="0.00", note="", bold=False, color=BLACK, size=10):
    ws.cell(row=row, column=2, value=label).font = Font(name=FONT, size=size, bold=bold)
    c = ws.cell(row=row, column=3, value=formula)
    c.font = Font(name=FONT, size=size, bold=True, color=color)
    c.number_format = fmt
    c.alignment = Alignment(horizontal="center")
    n = ws.cell(row=row, column=7, value=note)
    n.font = Font(name=FONT, size=9, italic=True, color="595959")
    n.alignment = Alignment(wrap_text=True, vertical="center")
    for cc in range(2, 8):
        ws.cell(row=row, column=cc).border = BORDER
    if note:
        ws.row_dimensions[row].height = rowh(note, 82, 15, 40)

M = f"'02 Indici'!$M${R0}:$M${LAST}"
Q = f"'02 Indici'!$Q${R0}:$Q${LAST}"
P = f"'02 Indici'!$P${R0}:$P${LAST}"
N = f"'02 Indici'!$N${R0}:$N${LAST}"
O = f"'02 Indici'!$O${R0}:$O${LAST}"

section(4, "A - PUNTEGGIO STRUMENTALE (32 indici, due fonti)")
line(5, "Indici valutati in totale", f"=COUNT({O})", "0", "Righe presenti nel foglio 02.")
line(6, "Indici normali", f'=COUNTIF({M},"Normale")', "0")
line(7, "Indici borderline (in zona di guardia)", f'=COUNTIF({M},"Borderline")', "0",
     "Dentro l'intervallo ma nell'ultimo 20% verso il lato critico. I referti li stampano come normali.")
line(8, "Indici fuori intervallo", f'=COUNTIF({M},"Fuori (alto)")+COUNTIF({M},"Fuori (basso)")', "0", "", bold=True, color="C00000")
line(9, "  di cui dall'articolo 1 (cardiologici)", f'=COUNTIFS({Q},"Art.1*",{M},"Fuori (alto)")+COUNTIFS({Q},"Art.1*",{M},"Fuori (basso)")', "0",
     "La fonte dichiara 1 indice fuori range su 16.")
line(10, "  di cui dall'articolo 2 (antropometrici)", f'=COUNTIFS({Q},"Art.2*",{M},"Fuori (alto)")+COUNTIFS({Q},"Art.2*",{M},"Fuori (basso)")', "0",
     "La fonte dichiara 6 indici oltre soglia su 12; applicando i cut-off delle sue stesse tabelle ne risultano 9. Vedi foglio 08.")
line(11, "Punteggio ponderato ottenuto", f"=SUM({P})", "0", "Somma di peso clinico x punteggio 0-3.")
line(12, "Punteggio massimo teorico", f"=SUM({N})*3", "0", "Se ogni indice pesato fosse gravemente fuori range.")
line(13, "COMPONENTE STRUMENTALE", "=IF(C12=0,0,C11/C12)", "0.0%", "Quota del massimo teorico raggiunta.", bold=True, color="C00000")

section(15, "B - FATTORI ANAMNESTICI CHE I REFERTI NON CONTENGONO")
hdr(ws, 16, ["", "Fattore", "Presente", "Peso", "Punti", "", "Evidenza / nota"], start=1, h=20)
FATT = [
 ("Fumo attivo (>= 20 sigarette al giorno)", f"=IF({I_CIG}>=20,1,0)", 3,
  "PMID 8689656 - assetto autonomico spostato verso il simpatico anche a distanza dall'ultima sigaretta."),
 ("Durata dell'abitudine tabagica >= 20 anni", f"=IF({I_YRS}>=20,1,0)", 2,
  "PMID 16029383 - la modulazione vagale cala al crescere degli anni, non solo delle sigarette."),
 ("Fenotipo TOFI (normopeso con obesita centrale)", f"=IF(AND({E_BMI}<25,{E_WHtR}>0.5),1,0)", 3,
  "PMID 26551006 - calcolato dagli indici: BMI sotto 25 con vita oltre meta dell'altezza. La mortalita cardiovascolare supera quella dell'obeso con distribuzione periferica."),
 ("Miosteatosi documentata", f"={I_MIO}", 2,
  "PMID 35227529 - infiltrazione adiposa del muscolo, non misurabile con l'antropometria di superficie. Un quadricipite infiltrato pesa uguale e lavora peggio."),
 ("Sospetta apnea ostruttiva del sonno mai indagata", f"={I_OSAS}", 3,
  "PMID 15301332 e 12609010 - il tono simpatico resta elevato anche in veglia. E il sospetto piu pesante proprio perche non e mai stato verificato."),
 ("Sedentarieta", f"={I_SED}", 1,
  "PMID 15793048 - l'esercizio supervisionato abbassa del 19% il prodotto cardiovascolare; la sua assenza toglie il principale fattore correttivo."),
 ("Incremento ponderale > 5 kg dall'eta di 30 anni", f"=IF({I_W}-{I_W30}>5,1,0)", 1,
  "Sei chili in vent'anni, trecento grammi l'anno: l'aumento piu banale del mondo. Sotto, circa dieci chili di grasso in piu e quattro di muscolo in meno."),
]
r = 17
f0 = r
for lab, presf, peso, ev in FATT:
    ws.cell(row=r, column=2, value=lab).font = Font(name=FONT, size=10)
    c = ws.cell(row=r, column=3, value=presf)
    c.font = Font(name=FONT, size=10, bold=True)
    c.alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=4, value=peso).font = Font(name=FONT, size=10, color=BLUE)
    ws.cell(row=r, column=4).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=5, value=f"=C{r}*D{r}").alignment = Alignment(horizontal="center")
    n = ws.cell(row=r, column=7, value=ev)
    n.font = Font(name=FONT, size=9, italic=True, color="595959")
    n.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = rowh(ev, 82, 26, 44)
    for cc in range(2, 8):
        ws.cell(row=r, column=cc).border = BORDER
        if ws.cell(row=r, column=cc).font.name != FONT:
            ws.cell(row=r, column=cc).font = Font(name=FONT, size=10)
    r += 1
f1 = r - 1
line(r, "Punti ottenuti", f"=SUMPRODUCT(C{f0}:C{f1},D{f0}:D{f1})", "0", "", bold=True)
line(r + 1, "Punti massimi", f"=SUM(D{f0}:D{f1})", "0", "", bold=True)
line(r + 2, "COMPONENTE ANAMNESTICA", f"=IF(C{r+1}=0,0,C{r}/C{r+1})", "0.0%",
     "La sarcopenia non compare qui perche e gia misurata fra gli indici del foglio 02, per evitare di contarla due volte.",
     bold=True, color="C00000")
FATT_PCT = f"C{r+2}"
FATT_GOT, FATT_MAX = f"C{r}", f"C{r+1}"

er = r + 4
section(er, "C - RISCHIO EPIDEMIOLOGICO CALCOLATO")
line(er + 1, "Scostamento dalla FC di minimo rischio (60-69 bpm)", f"={I_HR}-{I_HROPT}", "0", "bpm oltre il punto medio della fascia con mortalita minima.")
line(er + 2, "Scostamento dal limite superiore del referto", f"={I_HR}-{I_HRREF}", "0", "bpm oltre il valore che il referto considera normale.")
line(er + 3, "RR mortalita totale (+1,09 ogni 10 bpm)", f"=1.09^(({I_HR}-{I_HROPT})/10)", "0.000", "PMID 26598376. Stima di gradiente, non rischio assoluto individuale.")
line(er + 4, "RR mortalita cardiovascolare (+1,08 ogni 10 bpm)", f"=1.08^(({I_HR}-{I_HROPT})/10)", "0.000", "PMID 26598376.")
line(er + 5, "Soglia FC > 80 bpm superata", f'=IF({I_HR}>{I_HRTHR},"SI","NO")', "General",
     "Oltre gli 80 bpm il quadro cambia scala: RR 1,45 totale e 1,33 cardiovascolare.", color="C00000")
line(er + 6, "RR categoriale - mortalita totale", f'=IF({I_HR}>{I_HRTHR},1.45,"n.a.")', "0.00", "PMID 26598376.")
line(er + 7, "RR categoriale - mortalita cardiovascolare", f'=IF({I_HR}>{I_HRTHR},1.33,"n.a.")', "0.00", "PMID 26598376.")
line(er + 8, "HR se il valore >= 80 bpm si conferma nel tempo", f'=IF({I_HR}>={I_HRTHR},1.86,"n.a.")', "0.00",
     "PMID 28067310. Richiede misurazioni ripetute, oggi non disponibili.")
line(er + 9, "Battiti in eccesso all'anno", f"=({I_HR}-{I_HRREF})*60*24*365", "#,##0",
     "Il carico cumulativo che nessun sintomo segnala e nessuna visita registra.", bold=True)
line(er + 10, "Battiti in eccesso in 10 anni", f"=({I_HR}-{I_HRREF})*60*24*365*10", "#,##0")
line(er + 11, "Criterio vita per sindrome metabolica", f'=IF({I_WC}>={I_WCMS},"SODDISFATTO","NON soddisfatto")', "General",
     "Il criterio binario resta vuoto per 4 cm, mentre vita/altezza, conicita e ABSI sono gia tutti oltre i rispettivi cut-off. E il limite delle soglie a interruttore su variabili continue.")
line(er + 12, "Centimetri di vita che mancano al criterio", f"={I_WCMS}-{I_WC}", "0.0", "")

cr = er + 14
section(cr, "D - INDICE COMPOSITO DI PERICOLOSITA")
c = ws.cell(row=cr + 1, column=3, value=0.6)
line(cr + 1, "Peso della componente strumentale", 0.6, "0%", "Modificabile: le due componenti devono sommare a 1.")
ws.cell(row=cr + 1, column=3).font = Font(name=FONT, size=10, bold=True, color=BLUE)
ws.cell(row=cr + 1, column=3).fill = PatternFill("solid", fgColor=YEL)
line(cr + 2, "Peso della componente anamnestica", f"=1-C{cr+1}", "0%",
     "I referti coprono la maggior parte del giudizio, ma non tutto: il resto viene da cio che non misurano.")
line(cr + 3, "PUNTEGGIO DI PERICOLOSITA (0-100)", f"=(C13*C{cr+1}+{FATT_PCT}*C{cr+2})*100", "0.0",
     "Media pesata delle due componenti su scala centesimale.", bold=True, color="C00000", size=11)
ws.cell(row=cr + 3, column=3).font = Font(name=FONT, size=18, bold=True, color="C00000")
ws.row_dimensions[cr + 3].height = 30
line(cr + 4, "FASCIA", f'=IF(C{cr+3}<20,"BASSO",IF(C{cr+3}<40,"LIEVE",IF(C{cr+3}<60,"MODERATO",IF(C{cr+3}<80,"ELEVATO","ALTO"))))',
     "General", "Fasce: <20 basso | 20-40 lieve | 40-60 moderato | 60-80 elevato | >=80 alto.", bold=True, size=12)
ws.cell(row=cr + 4, column=3).font = Font(name=FONT, size=15, bold=True, color="C00000")
ws.row_dimensions[cr + 4].height = 26

ws.cell(row=cr + 6, column=2, value="Lettura sintetica").font = Font(name=FONT, size=10, bold=True)
c = ws.cell(row=cr + 6, column=3, value=(
    f'="Su "&C5&" indici ricalcolati, "&C8&" risultano fuori intervallo nella direzione critica e "&C7&'
    f'" sono in zona di guardia pur essendo stampati come normali. "'
    f'&IF({I_HR}>{I_HRTHR},"La frequenza a riposo supera la soglia degli 80 bpm, oltre la quale il rischio cambia scala. ","")'
    f'&IF(AND({E_BMI}<25,{E_WHtR}>0.5),"Il BMI resta sotto 25 mentre la vita supera la meta dell altezza: e il fenotipo che il peso corporeo nasconde. ","")'
    f'&"Fattori anamnestici presenti: "&{FATT_GOT}&" punti su "&{FATT_MAX}&". "'
    f'&"Accertamenti mancanti che cambierebbero il giudizio: test da sforzo, studio del sonno e misurazioni ripetute della frequenza."'))
c.font = Font(name=FONT, size=10)
c.alignment = Alignment(wrap_text=True, vertical="top")
c.fill = WARN_FILL
ws.merge_cells(start_row=cr + 6, start_column=3, end_row=cr + 9, end_column=7)
for cc in range(3, 8):
    for rr in range(cr + 6, cr + 10):
        ws.cell(row=rr, column=cc).fill = WARN_FILL
        ws.cell(row=rr, column=cc).border = BORDER

mr = cr + 11
section(mr, "E - COSA I REFERTI NON POSSONO DIRE")
MANCA = [
 ("Test da sforzo", "Senza di esso non e valutabile la riserva effettivamente impegnata, il parametro con HR 2,8 per mortalita cardiovascolare (PMID 17446799)."),
 ("Studio del sonno", "Russamento e pause respiratorie riferiti ma mai indagati. Le apnee spiegherebbero una quota della frequenza elevata (PMID 15301332)."),
 ("Misurazioni ripetute della FC a riposo", "L'HR 1,86 di Kailuan si applica alla persistenza del valore su piu misurazioni, non a un singolo prelievo (PMID 28067310)."),
 ("Variabilita della frequenza cardiaca (HRV)", "LF, HF e guadagno barocettivo sono i parametri su cui il fumo agisce, e nessuno di essi compare nei referti."),
 ("Glicemia, insulinemia, enzimi epatici", "VAI e CMI segnalano resistenza insulinica e disfunzione del tessuto adiposo, ma la conferma diretta richiede i marcatori di sensibilita insulinica e di steatosi."),
 ("DEXA o bioimpedenziometria", "Gli indici antropometrici stratificano il rischio a costo quasi nullo, ma non quantificano la composizione corporea: sarcopenia e miosteatosi restano stime."),
]
r = mr + 1
for lab, why in MANCA:
    ws.cell(row=r, column=2, value=lab).font = Font(name=FONT, size=10, bold=True)
    c = ws.cell(row=r, column=3, value=why)
    c.font = Font(name=FONT, size=9)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
    ws.row_dimensions[r].height = 30
    for cc in range(2, 8):
        ws.cell(row=r, column=cc).border = BORDER
    r += 1

dr = r + 1
ws.cell(row=dr, column=2, value="AVVERTENZA").font = Font(name=FONT, size=10, bold=True, color="C00000")
c = ws.cell(row=dr, column=3, value=(
    "Il punteggio composito e una sintesi documentale costruita su pesi scelti dall'analista, non uno score clinico "
    "validato. Serve a rendere esplicito e discutibile un giudizio, non a sostituirlo. I pesi sono modificabili nella "
    "colonna 'Peso' del foglio 02 e nella colonna D della sezione B. Prima di usare i numeri, leggere il foglio "
    "'08 Controlli': sette valori pubblicati non si riproducono dai dati grezzi."))
c.font = Font(name=FONT, size=9, italic=True)
c.alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells(start_row=dr, start_column=3, end_row=dr + 2, end_column=7)
for cc in range(2, 8):
    for rr in range(dr, dr + 3):
        ws.cell(row=rr, column=cc).fill = WARN_FILL
        ws.cell(row=rr, column=cc).border = BORDER

# =====================================================================
# 07 SCENARI
# =====================================================================
ws = wb.create_sheet("07 Scenari")
setw(ws, {"A": 4, "B": 38, "C": 10, "D": 10, "E": 10, "F": 11, "G": 10, "H": 10,
          "I": 13, "J": 14, "K": 10, "L": 50})
title(ws, "07 - Scenari di intervento",
      "Le celle gialle sono le leve dello scenario; il resto e calcolato. I valori sono proiezioni sull'ordine di grandezza degli effetti riportati dagli studi, non risultati misurati.")

ws.cell(row=4, column=2, value="TAVOLA 1 - Leve cardiovascolari").font = Font(name=FONT, size=11, bold=True, color="1F3864")
hdr(ws, 5, ["", "Scenario", "FC\n(bpm)", "SV\n(mL)", "SBP\n(mmHg)", "MAP\n(mmHg)", "CO\n(L/min)",
            "MSI", "RPP\n(bpm*mmHg)", "Battiti extra\n/anno", "RR\nmort.tot.", "Commento"], start=1)
SC = [
 ("Situazione attuale", 82, 66, 128, "Quindici indici cardiologici su sedici stampati come normali."),
 ("Esercizio supervisionato, 3 mesi", 72, 66, 124, "PMID 15793048: -19% sul prodotto cardiovascolare. Nel gruppo non supervisionato il valore era rimasto invariato."),
 ("Cessazione del fumo, oltre 4 settimane", 77, 66, 126, "PMID 8651120: la FC cala e la variabilita migliora, ma a 4 settimane resta sopra i valori attesi per l'eta."),
 ("Esercizio + cessazione", 70, 70, 122, "Le due leve insieme, con recupero parziale di gittata dovuto al condizionamento."),
 ("Pompa piu capiente a pari portata", 68, 80, 128, "Stessa portata ottenuta con gittata nel medio dell'intervallo: il confronto che mostra cosa nasconde una portata normale."),
 ("Obiettivo fascia di minimo rischio", 65, 80, 120, "Frequenza nella fascia 60-69 bpm, minimo di mortalita secondo PMID 11337213."),
]
r = 6
s0 = r
for i, (nome, hr, sv, sbp, comm) in enumerate(SC, 1):
    ws.cell(row=r, column=1, value=i).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=2, value=nome).font = Font(name=FONT, size=10, bold=(i == 1))
    for col, val in ((3, hr), (4, sv), (5, sbp)):
        c = ws.cell(row=r, column=col, value=val)
        c.font = Font(name=FONT, size=10, bold=True, color=BLUE)
        c.fill = PatternFill("solid", fgColor=YEL)
    ws.cell(row=r, column=6, value=f"={I_DBP}+0.4*(E{r}-{I_DBP})").number_format = "0"
    ws.cell(row=r, column=7, value=f"=D{r}*C{r}/1000").number_format = "0.0"
    ws.cell(row=r, column=8, value=f"=C{r}/F{r}").number_format = "0.00"
    ws.cell(row=r, column=9, value=f"=C{r}*E{r}").number_format = "#,##0"
    ws.cell(row=r, column=10, value=f"=MAX(0,C{r}-{I_HRREF})*60*24*365").number_format = "#,##0"
    ws.cell(row=r, column=11, value=f"=1.09^((C{r}-{I_HROPT})/10)").number_format = "0.000"
    ws.cell(row=r, column=12, value=comm).font = Font(name=FONT, size=9)
    for cc in range(1, 13):
        cell = ws.cell(row=r, column=cc)
        cell.border = BORDER
        if cell.font.name != FONT:
            cell.font = Font(name=FONT, size=10)
        cell.alignment = Alignment(horizontal="center" if 3 <= cc <= 11 or cc == 1 else "left",
                                   vertical="center", wrap_text=(cc in (2, 12)))
    ws.row_dimensions[r].height = rowh(comm, 56, 28, 52)
    r += 1
s1 = r - 1

r += 1
ws.cell(row=r, column=2, value="TAVOLA 2 - Leve antropometriche").font = Font(name=FONT, size=11, bold=True, color="1F3864")
hdr(ws, r + 1, ["", "Scenario", "Peso\n(kg)", "Vita\n(cm)", "Fianchi\n(cm)", "BMI", "WHtR", "Conicity",
                "LAP", "CMI", "VAI", "Commento"], start=1)
SC2 = [
 ("Situazione attuale", 74.0, 97, 100, "BMI normale, sei indici di forma e ibridi oltre soglia."),
 ("-5 cm di vita a pari peso", 74.0, 92, 99, "Ricomposizione senza calo ponderale: e cio che l'esercizio contro resistenza produce per primo."),
 ("-3 kg con -6 cm di vita", 71.0, 91, 98, "Rientro nella forbice del peso desiderabile (68-71 kg secondo Lorenz e Creff)."),
 ("Vita sotto meta altezza (WHtR 0,50)", 71.0, 87, 97, "Soglia della regola aurea: la vita misura meno della meta dell'altezza."),
 ("Obiettivo pieno su tutti gli indici di forma", 69.0, 84, 96, "Rientro sotto il cut-off di 0,52 per WHtR e sotto 1,25 per il Conicity Index."),
]
r2 = r + 2
t0 = r2
for i, (nome, w, wc, hip, comm) in enumerate(SC2, 1):
    ws.cell(row=r2, column=1, value=i).alignment = Alignment(horizontal="center")
    ws.cell(row=r2, column=2, value=nome).font = Font(name=FONT, size=10, bold=(i == 1))
    for col, val in ((3, w), (4, wc), (5, hip)):
        c = ws.cell(row=r2, column=col, value=val)
        c.font = Font(name=FONT, size=10, bold=True, color=BLUE)
        c.fill = PatternFill("solid", fgColor=YEL)
    ws.cell(row=r2, column=6, value=f"=C{r2}/({I_H}/100)^2").number_format = "0.00"
    ws.cell(row=r2, column=7, value=f"=D{r2}/{I_H}").number_format = "0.000"
    ws.cell(row=r2, column=8, value=f"=(D{r2}/100)/(0.109*SQRT(C{r2}/({I_H}/100)))").number_format = "0.00"
    ws.cell(row=r2, column=9, value=f"=(D{r2}-65)*({I_TG}/88.57)").number_format = "0.0"
    ws.cell(row=r2, column=10, value=f"=(D{r2}/{I_H})*(({I_TG}/88.57)/({I_HDL}/38.67))").number_format = "0.00"
    ws.cell(row=r2, column=11, value=f"=(D{r2}/(39.68+1.88*F{r2}))*(({I_TG}/88.57)/1.03)*(1.31/({I_HDL}/38.67))").number_format = "0.00"
    ws.cell(row=r2, column=12, value=comm).font = Font(name=FONT, size=9)
    for cc in range(1, 13):
        cell = ws.cell(row=r2, column=cc)
        cell.border = BORDER
        if cell.font.name != FONT:
            cell.font = Font(name=FONT, size=10)
        cell.alignment = Alignment(horizontal="center" if 3 <= cc <= 11 or cc == 1 else "left",
                                   vertical="center", wrap_text=(cc in (2, 12)))
    ws.row_dimensions[r2].height = rowh(comm, 56, 28, 52)
    r2 += 1
t1 = r2 - 1

nr = r2 + 1
ws.cell(row=nr, column=2, value="Nota: i cut-off da rientrare sono WHtR 0,52 - Conicity 1,25 - LAP 26,7 - CMI 0,39 - VAI 1,92. "
                                "Gli indici ibridi migliorano anche solo riducendo la vita, perche la vita e il loro primo fattore.").font = Font(name=FONT, size=9, italic=True, color="595959")
ws.merge_cells(start_row=nr, start_column=2, end_row=nr + 1, end_column=12)
ws.cell(row=nr, column=2).alignment = Alignment(wrap_text=True, vertical="top")

# =====================================================================
# 08 CONTROLLI
# =====================================================================
ws = wb.create_sheet("08 Controlli")
setw(ws, {"A": 5, "B": 11, "C": 38, "D": 17, "E": 19, "F": 20, "G": 74})
title(ws, "08 - Controlli di riproducibilita",
      "Ogni valore pubblicato dalle fonti e stato ricalcolato dai dati grezzi. Qui e registrato cosa torna e cosa no.")
hdr(ws, 4, ["#", "Fonte", "Elemento verificato", "Valore pubblicato", "Valore ricalcolato", "Esito", "Impatto sul giudizio"])

CTRL = [
 (A2, "BMI da peso e altezza", "24,09", "24,44", "Scostamento",
  "74 kg su 174 cm danno 24,44. Per ottenere 24,09 servirebbero 72,9 kg a 174 cm, oppure 175,3 cm a 74 kg. Non e un arrotondamento: uno dei tre numeri pubblicati e incoerente con gli altri due. Il BMI entra nel VAI, nell'ABSI e nella correzione del polpaccio, quindi l'errore si propaga su tre indici."),
 (A2, "BMI usato internamente dal referto", "-", "24,09", "Dedotto",
  "Il VAI pubblicato (2,29) si riproduce solo con BMI 24,09; con 24,44 si otterrebbe 2,27. Il referto usa dunque internamente il BMI dichiarato, non quello ricavabile da peso e altezza. Il modello mantiene entrambi i valori senza sceglierne uno."),
 (A2, "Abdominal Volume Index", "11,89", "18,82", "Non riproducibile",
  "Scostamento del 58%. La formula standard (2 x vita^2 + 0,7 x (vita-fianchi)^2)/1000 non restituisce il valore pubblicato, e non e stata trovata alcuna variante che lo faccia. L'esito qualitativo non cambia: l'indice resta dentro il riferimento con entrambi i valori."),
 (A2, "Percentuale di grasso corporeo", "23%", "24,6% (Deurenberg)", "Non riproducibile",
  "La formula non e dichiarata. Deurenberg e la candidata piu probabile ma restituisce 24,6% con il BMI ricalcolato e 24,2% con quello del referto. Scostamento di circa 1,2 punti. Poco rilevante: e comunque una stima derivata dal BMI."),
 (A2, "Circonferenza polpaccio corretta per BMI", "33,5 cm", "34,0 cm", "Non riproducibile",
  "IMPATTO ALTO. La correzione categoriale del consenso EWGSOP2 non prevede aggiustamenti per BMI fra 18,5 e 24,9, quindi il valore resterebbe 34,0 cm, cioe esattamente sulla soglia invece che sotto. La diagnosi di sarcopenia moderata dipende interamente da una correzione di 5 millimetri la cui formula non e pubblicata."),
 (A2, "Conteggio degli indici oltre soglia", "6 su 12", "9 su 12", "Scostamento",
  "Applicando i cut-off riportati dalle tabelle della fonte stessa risultano oltre soglia: vita/altezza, vita/fianchi, conicita, BAI, LAP, CMI, VAI, ABSI e polpaccio. Restano dentro solo percentuale di grasso, BRI e AVI. La direzione della conclusione non cambia, anzi si rafforza."),
 (A2, "ABSI z-score", "1,52", "non ricalcolabile", "Non verificabile",
  "Il valore grezzo di ABSI si ricalcola (0,0873), ma lo z-score richiede le tabelle NHANES per eta e sesso, che la fonte non riporta. Nel foglio 02 il valore viene inserito a mano e non e verificato."),
 (A1, "Numero di indici cardiologici", "16", "15 tabulati", "Scostamento",
  "La fonte dichiara sedici indici ma le tre tabelle ne contengono quindici. Il modello include la superficie corporea come sedicesima riga informativa, essendo l'unico altro valore numerico del referto citato nel testo."),
 (A1, "Pressione arteriosa media", "101 mmHg", "98,7 (classica) / 101,6 (DBP+0,4xPP)", "Formula ricostruita",
  "La formula classica non riproduce il valore pubblicato. Rilevante perche la MAP e il denominatore del Modified Shock Index: con la formula classica il MSI passerebbe da 0,81 a 0,83."),
 (A1, "Left Cardiac Work Index", "627 mW/m2", "639 mW/m2", "Scostamento minore",
  "Circa il 2%, interamente dovuto agli arrotondamenti a monte: usando il Cardiac Index arrotondato a 2,8 si ottengono 628. Nessun impatto sul giudizio."),
 (A1, "Resistenze vascolari sistemiche", "2,7 kU*m2", "2,70 con CVP assunta a 5 mmHg", "Riproducibile con assunzione",
  "La pressione venosa centrale non e pubblicata. Con CVP 5 mmHg il calcolo coincide, con CVP 0 si otterrebbe 2,8, che spingerebbe l'indice sopra il limite superiore di 2,8. L'assunzione e quindi tutt'altro che neutra."),
 (A1, "Indice di rigidita arteriosa", "0,44", "0,436 (PP/MAP)", "Formula ricostruita",
  "Riproduce il valore pubblicato. Il rapporto alternativo PP/SBP darebbe 0,34."),
 (A1, "Zone di frequenza cardiaca", "87 / 105 / 122 / 140 / 157", "identiche con troncamento", "Riproducibile",
  "Le soglie si riproducono esattamente troncando all'intero le percentuali della FC massima, non arrotondando."),
 ("Entrambe", "Tutti gli altri indici", "-", "coincidono", "Riproducibile",
  "Riproducono il valore pubblicato entro l'arrotondamento: pressione di pulsazione, riserva, Modified Shock Index, prodotto cardiovascolare, portata, Cardiac Index, morfotipo, pesi ideali di Lorenz e Creff, vita/altezza, vita/fianchi, Conicity Index, BAI, BRI, LAP, CMI e VAI."),
]
r = 5
for i, (fonte, elem, pub, calc, esito, imp) in enumerate(CTRL, 1):
    ws.cell(row=r, column=1, value=i).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(row=r, column=2, value=fonte.split(" - ")[0]).font = Font(name=FONT, size=9, color="595959")
    ws.cell(row=r, column=3, value=elem).font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=r, column=4, value=pub).font = Font(name=FONT, size=10, color=BLUE)
    ws.cell(row=r, column=5, value=calc).font = Font(name=FONT, size=10, bold=True)
    bad = esito in ("Non riproducibile", "Scostamento", "Non verificabile")
    ws.cell(row=r, column=6, value=esito).font = Font(name=FONT, size=10, bold=True,
                                                     color="C00000" if bad else "008000")
    ws.cell(row=r, column=7, value=imp).font = Font(name=FONT, size=9)
    ws.row_dimensions[r].height = rowh(imp, 90, 30, 84)
    for cc in range(1, 8):
        cell = ws.cell(row=r, column=cc)
        cell.border = BORDER
        if cell.font.name != FONT:
            cell.font = Font(name=FONT, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top",
                                   horizontal="center" if cc in (1, 2, 4, 5, 6) else "left")
        if bad:
            cell.fill = ERR_FILL
    r += 1

sr = r + 1
ws.cell(row=sr, column=3, value="Sintesi").font = Font(name=FONT, size=10, bold=True)
for k, (lab, val) in enumerate([
        ("Elementi verificati", f"=COUNTA(C5:C{r-1})"),
        ("Riproducibili senza riserve", f'=COUNTIF(F5:F{r-1},"Riproducibile")'),
        ("Formule ricostruite dall'analista", f'=COUNTIF(F5:F{r-1},"Formula ricostruita")+COUNTIF(F5:F{r-1},"Riproducibile con assunzione")'),
        ("Scostamenti o valori non riproducibili", f'=COUNTIF(F5:F{r-1},"Non riproducibile")+COUNTIF(F5:F{r-1},"Scostamento")+COUNTIF(F5:F{r-1},"Non verificabile")')]):
    ws.cell(row=sr + 1 + k, column=3, value=lab).font = Font(name=FONT, size=10)
    c = ws.cell(row=sr + 1 + k, column=5, value=val)
    c.font = Font(name=FONT, size=11, bold=True, color="C00000" if k == 3 else BLACK)
    c.alignment = Alignment(horizontal="center")
    for cc in range(3, 7):
        ws.cell(row=sr + 1 + k, column=cc).border = BORDER

cr2 = sr + 6
ws.cell(row=cr2, column=3, value="Come leggere questo foglio").font = Font(name=FONT, size=10, bold=True, color="C00000")
c = ws.cell(row=cr2 + 1, column=3, value=(
    "Nessuno di questi scostamenti ribalta le conclusioni delle fonti: gli indici che segnalano rischio lo segnalano "
    "con entrambi i valori, e il conteggio degli indici oltre soglia risulta piu alto del dichiarato, non piu basso. "
    "Due elementi meritano pero attenzione. Il primo e la diagnosi di sarcopenia, che poggia su una correzione di "
    "cinque millimetri la cui formula non e pubblicata e che con il criterio del consenso europeo non verrebbe "
    "applicata. Il secondo e il BMI, che non discende dal peso e dall'altezza dichiarati e che alimenta tre indici "
    "derivati. Entrambi si risolvono chiedendo alla fonte i dati grezzi o la formula usata."))
c.font = Font(name=FONT, size=9, italic=True)
c.alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells(start_row=cr2 + 1, start_column=3, end_row=cr2 + 4, end_column=7)
for cc in range(3, 8):
    for rr in range(cr2 + 1, cr2 + 5):
        ws.cell(row=rr, column=cc).fill = WARN_FILL
        ws.cell(row=rr, column=cc).border = BORDER

# =====================================================================
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None and (cell.font is None or cell.font.name != FONT):
                f = cell.font
                cell.font = Font(name=FONT, size=f.size or 10, bold=f.bold, italic=f.italic,
                                 color=f.color, underline=f.underline)
    sheet.sheet_view.showGridLines = False

wb.save(OUT)
print("scritto:", OUT, "| righe indici:", R0, "-", LAST)
